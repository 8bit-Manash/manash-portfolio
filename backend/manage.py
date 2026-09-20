"""
Manash Madhukar Portfolio — Django Backend
Saves contact form submissions to contact_submissions.xlsx
AND emails a notification to manastiwary2067@gmail.com on every submission.

Run:
    cd backend
    pip install -r requirements.txt

    # Required for email to actually send (Gmail SMTP):
    # Sending FROM: manas20tiwary@gmail.com  (already set as the default below)
    # 1. Turn on 2-Step Verification on that Gmail account.
    # 2. Create an "App Password": https://myaccount.google.com/apppasswords
    # 3. Set the password either way (both work):
    #      a) shell:  export EMAIL_HOST_PASSWORD="the16charapppassword"
    #                 (Windows: set EMAIL_HOST_PASSWORD=...)
    #      b) file:   create backend/.env containing:
    #                   EMAIL_HOST_PASSWORD="the16charapppassword"
    #                   EMAIL_HOST_USER="manas20tiwary@gmail.com"
    #    To send from a different address instead, also set EMAIL_HOST_USER.
    #    Without EMAIL_HOST_PASSWORD set, the form still works and still saves
    #    to Excel — it just skips sending the email and logs a warning instead.

    python manage.py migrate
    python manage.py runserver 8000
"""
import os, sys, json, pathlib
from datetime import datetime

BASE_DIR = pathlib.Path(__file__).resolve().parent

# ── Load a .env file sitting next to this file, if one exists.
#    This lets EMAIL_HOST_USER / EMAIL_HOST_PASSWORD be set either as a real
#    shell env var OR just saved in backend/.env — both work the same way. ──
def _load_dotenv(path):
    if not path.exists():
        return
    for line in path.read_text().splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, value = line.partition("=")
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        os.environ.setdefault(key, value)

_load_dotenv(BASE_DIR / ".env")

import django
from django.conf import settings

EXCEL_FILE = BASE_DIR / "contact_submissions.xlsx"

# ── Where notification emails get sent — your inbox ──
NOTIFY_EMAIL = "manastiwary2067@gmail.com"

# ── SMTP credentials come from environment variables, never hardcoded here.
#    Set these before running the server (see README instructions below):
#      EMAIL_HOST_USER      -> the Gmail address that SENDS the notification
#      EMAIL_HOST_PASSWORD  -> a 16-character Gmail "App Password" (not your normal password)
EMAIL_HOST_USER = os.environ.get("EMAIL_HOST_USER", "manas20tiwary@gmail.com")
EMAIL_HOST_PASSWORD = os.environ.get("EMAIL_HOST_PASSWORD", "")

# ── Fix for macOS "CERTIFICATE_VERIFY_FAILED" when sending email ──
#    Some Python installs on Mac (and some Linux setups) don't have access to
#    the system's root certificates, so any TLS connection — including Gmail
#    SMTP — fails to verify. Pointing Python at certifi's bundled certs fixes
#    this permanently without needing to set anything in the terminal.
try:
    import certifi
    os.environ.setdefault("SSL_CERT_FILE", certifi.where())
except ImportError:
    pass  # certifi not installed — run: pip install certifi

# ── Django config ─────────────────────────────────────────────────────────────
if not settings.configured:
    settings.configure(
        DEBUG=True,
        SECRET_KEY="manash-portfolio-change-this-in-production-abc123xyz",
        ALLOWED_HOSTS=["*"],
        INSTALLED_APPS=[
            "django.contrib.contenttypes",
            "django.contrib.auth",
            "corsheaders",
        ],
        MIDDLEWARE=[
            "corsheaders.middleware.CorsMiddleware",
            "django.middleware.common.CommonMiddleware",
        ],
        CORS_ALLOW_ALL_ORIGINS=True,       # Lock this down in production
        ROOT_URLCONF=__name__,
        DATABASES={
            "default": {
                "ENGINE": "django.db.backends.sqlite3",
                "NAME": BASE_DIR / "db.sqlite3",
            }
        },
        # ── Email (Gmail SMTP) ──
        EMAIL_BACKEND="django.core.mail.backends.smtp.EmailBackend",
        EMAIL_HOST="smtp.gmail.com",
        EMAIL_PORT=587,
        EMAIL_USE_TLS=True,
        EMAIL_HOST_USER=EMAIL_HOST_USER,
        EMAIL_HOST_PASSWORD=EMAIL_HOST_PASSWORD,
        DEFAULT_FROM_EMAIL=EMAIL_HOST_USER,
    )

django.setup()

# ── Excel helpers ──────────────────────────────────────────────────────────────
def get_or_create_workbook():
    try:
        from openpyxl import load_workbook
        wb = load_workbook(EXCEL_FILE)
        ws = wb.active
    except FileNotFoundError:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook()
        ws = wb.active
        ws.title = "Contact Submissions"

        headers = ["#", "Name", "Email", "Message", "Timestamp", "IP"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(color="00F5FF", bold=True, name="Consolas")
            cell.fill = PatternFill(start_color="0A0A12", end_color="0A0A12", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        ws.column_dimensions["A"].width = 5
        ws.column_dimensions["B"].width = 25
        ws.column_dimensions["C"].width = 35
        ws.column_dimensions["D"].width = 60
        ws.column_dimensions["E"].width = 22
        ws.column_dimensions["F"].width = 18
        wb.save(EXCEL_FILE)
    return wb, ws


def save_to_excel(name, email, message, ip=""):
    wb, ws = get_or_create_workbook()
    from openpyxl.styles import Font, Alignment
    next_row = ws.max_row + 1
    data = [next_row - 1, name, email, message,
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"), ip]
    for col, value in enumerate(data, 1):
        cell = ws.cell(row=next_row, column=col, value=value)
        cell.font = Font(name="Calibri", size=11)
        if col == 4:
            cell.alignment = Alignment(wrap_text=True)
    wb.save(EXCEL_FILE)
    print(f"[Excel] Saved row {next_row}: {name} <{email}>")
    return next_row - 1


# ── Email helper ───────────────────────────────────────────────────────────────
def send_notification_email(name, email, message):
    """Send a notification email to NOTIFY_EMAIL whenever the contact form is
    submitted. Returns True/False — never raises, so a missing/bad SMTP config
    never breaks the form submission itself (the Excel save already succeeded)."""
    if not EMAIL_HOST_USER or not EMAIL_HOST_PASSWORD:
        print("[EMAIL] Skipped — EMAIL_HOST_USER/EMAIL_HOST_PASSWORD not set in environment.")
        return False
    try:
        from django.core.mail import send_mail
        subject = f"New Portfolio Contact — {name}"
        body = (
            f"You've got a new message from your portfolio contact form.\n\n"
            f"Name:    {name}\n"
            f"Email:   {email}\n"
            f"Message:\n{message}\n\n"
            f"— Sent automatically from manashmadhukar.dev"
        )
        send_mail(
            subject=subject,
            message=body,
            from_email=EMAIL_HOST_USER,
            recipient_list=[NOTIFY_EMAIL],
            fail_silently=False,
        )
        print(f"[EMAIL] Notification sent to {NOTIFY_EMAIL} for submission from {email}")
        return True
    except Exception as e:
        print(f"[EMAIL ERROR] {e}")
        return False


# ── Views ──────────────────────────────────────────────────────────────────────
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods


@csrf_exempt
@require_http_methods(["POST", "OPTIONS"])
def contact_view(request):
    if request.method == "OPTIONS":
        r = JsonResponse({})
        r["Access-Control-Allow-Origin"] = "*"
        r["Access-Control-Allow-Headers"] = "Content-Type"
        return r
    try:
        body = json.loads(request.body)
        name    = body.get("name", "").strip()
        email   = body.get("email", "").strip()
        message = body.get("message", "").strip()

        errors = {}
        if not name:               errors["name"] = "Name is required"
        if not email or "@" not in email: errors["email"] = "Valid email is required"
        if not message:            errors["message"] = "Message is required"
        if errors:
            return JsonResponse({"success": False, "errors": errors}, status=400)

        ip = request.META.get("HTTP_X_FORWARDED_FOR", request.META.get("REMOTE_ADDR", ""))
        row = save_to_excel(name, email, message, ip)

        # ── Send notification email — failure here must NOT fail the request,
        #    the submission is already safely saved to Excel above. ──
        email_sent = send_notification_email(name, email, message)
        if not email_sent:
            print("[EMAIL] Notification email failed to send — check EMAIL_HOST_USER / EMAIL_HOST_PASSWORD env vars.")

        return JsonResponse({
            "success": True,
            "message": "Your connection request has been submitted. I'll connect with you soon!",
            "id": row,
        }, status=201)

    except json.JSONDecodeError:
        return JsonResponse({"success": False, "error": "Invalid JSON"}, status=400)
    except Exception as e:
        print(f"[ERROR] {e}")
        return JsonResponse({"success": False, "error": "Server error"}, status=500)


@csrf_exempt
@require_http_methods(["GET"])
def submissions_view(request):
    """GET /api/submissions/ — view all saved contacts (admin only in prod)"""
    try:
        wb, ws = get_or_create_workbook()
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                rows.append({"id": row[0], "name": row[1], "email": row[2],
                             "message": row[3], "timestamp": str(row[4]), "ip": row[5]})
        return JsonResponse({"success": True, "count": len(rows), "submissions": rows})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)}, status=500)


def health_view(request):
    count = 0
    if EXCEL_FILE.exists():
        try:
            _, ws = get_or_create_workbook()
            count = ws.max_row - 1
        except:
            pass
    return JsonResponse({"status": "ok", "service": "Manash Portfolio Backend", "submissions": count})


# ── URLs ───────────────────────────────────────────────────────────────────────
from django.urls import path

urlpatterns = [
    path("api/contact/",     contact_view),
    path("api/submissions/", submissions_view),
    path("api/health/",      health_view),
]

# ── WSGI app for production servers (gunicorn on Render uses this) ─────────────
from django.core.wsgi import get_wsgi_application
application = get_wsgi_application()

# ── Entry point ────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    from django.core.management import execute_from_command_line
    execute_from_command_line(sys.argv if len(sys.argv) > 1 else [__file__, "runserver", "0.0.0.0:8000"])